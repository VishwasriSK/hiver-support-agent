import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation


INPUT_PATH = "golden_set/golden_200_to_label.csv"
OUTPUT_PATH = "golden_set/golden_200_to_label.xlsx"


# --------------------------------------------------
# Load CSV
# --------------------------------------------------

print("Loading golden labeling CSV...")

df = pd.read_csv(INPUT_PATH)

print("Rows:", len(df))


# --------------------------------------------------
# Create Excel file
# --------------------------------------------------

df.to_excel(
    OUTPUT_PATH,
    index=False,
    sheet_name="Golden Set"
)


# --------------------------------------------------
# Open workbook
# --------------------------------------------------

wb = load_workbook(OUTPUT_PATH)

ws = wb["Golden Set"]


# --------------------------------------------------
# Header formatting
# --------------------------------------------------

for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


# --------------------------------------------------
# Wrap text
# --------------------------------------------------

for row in ws.iter_rows():

    for cell in row:

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )


# --------------------------------------------------
# Column widths
# --------------------------------------------------

widths = {
    "A": 10,   # gold_id
    "B": 20,   # customer_tweet_id
    "C": 65,   # customer_text
    "D": 65,   # brand_reply
    "E": 32,   # candidate_intent
    "F": 32,   # gold_intent
    "G": 40    # review_notes
}


for column, width in widths.items():

    ws.column_dimensions[column].width = width


# --------------------------------------------------
# Row height
# --------------------------------------------------

ws.row_dimensions[1].height = 30

for row in range(2, ws.max_row + 1):
    ws.row_dimensions[row].height = 80


# --------------------------------------------------
# Freeze header
# --------------------------------------------------

ws.freeze_panes = "A2"


# --------------------------------------------------
# Add filter
# --------------------------------------------------

ws.auto_filter.ref = ws.dimensions


# --------------------------------------------------
# Gold intent dropdown
# --------------------------------------------------

intents = [
    "Delivery & Tracking",
    "Order Management",
    "Payment & Charges",
    "Returns & Refunds",
    "Account & Login",
    "Prime Membership & Benefits",
    "Digital Services & Devices",
    "Product Problems"
]


intent_string = ",".join(intents)


validation = DataValidation(
    type="list",
    formula1=f'"{intent_string}"',
    allow_blank=False
)

validation.error = "Please select a valid intent."
validation.errorTitle = "Invalid Intent"

validation.prompt = "Select the correct gold intent."
validation.promptTitle = "Gold Intent"

ws.add_data_validation(validation)


# Apply dropdown to gold_intent column

validation.add(
    f"F2:F{ws.max_row}"
)


# --------------------------------------------------
# Save
# --------------------------------------------------

wb.save(OUTPUT_PATH)


print("\n--------------------------------")
print("Excel labeling file created!")
print("--------------------------------")

print("Rows:", len(df))

print("Saved to:")
print(OUTPUT_PATH)